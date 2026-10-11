from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="1F3A5F", end_color="1F3A5F", fill_type="solid")
BODY_FONT = Font(name="Arial")
TITLE_FONT = Font(name="Arial", bold=True, size=13)

wb = Workbook()
wb.remove(wb.active)

# --- Sheet 1: Pivot Notes ---
notes = wb.create_sheet("Pivot Notes")
notes["A1"] = "Pivot Notes"
notes["A1"].font = TITLE_FONT
notes["A3"] = (
    "Import checkout_sessions_export.csv, insert a pivot: "
    "rows = last_step_reached, columns = input_mode, values = sum of cart_value, "
    "filter converted = N."
)
notes["A3"].font = BODY_FONT
notes["A3"].alignment = Alignment(wrap_text=True, vertical="top")
notes.column_dimensions["A"].width = 90
notes.row_dimensions[3].height = 60

# --- Sheet 2: Defect Log ---
log = wb.create_sheet("Defect Log")
headers = [
    "ID",
    "Checkout step",
    "Input mode affected",
    "Component",
    "Defect",
    "Revenue at risk",
    "Est. hours",
    "Revenue at risk per hour",
    "Rank",
    "Status",
    "Notes",
]
for col_index, header in enumerate(headers, start=1):
    cell = log.cell(row=1, column=col_index, value=header)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center")
log.freeze_panes = "A2"

column_widths = [6, 16, 18, 22, 32, 16, 12, 22, 8, 12, 28]
for col_index, width in enumerate(column_widths, start=1):
    log.column_dimensions[get_column_letter(col_index)].width = width

for row in range(2, 8):
    log.cell(row=row, column=8, value=f'=IF(G{row}>0,F{row}/G{row},"")').font = BODY_FONT
    log.cell(row=row, column=9, value=f'=IF(H{row}="","",RANK(H{row},$H$2:$H$7))').font = BODY_FONT
    for col_index in range(1, 12):
        if col_index in (8, 9):
            continue
        log.cell(row=row, column=col_index).font = BODY_FONT

status_validation = DataValidation(
    type="list", formula1='"Open,Fixed,Backlog"', allow_blank=True, showDropDown=False
)
log.add_data_validation(status_validation)
status_validation.add("J2:J7")

# --- Sheet 3: Summary ---
summary = wb.create_sheet("Summary")
summary["A1"] = "Summary"
summary["A1"].font = TITLE_FONT
labels = [
    "Total revenue at risk (before)",
    "Revenue at risk addressed by fixes",
    "Revenue at risk remaining",
    "Release recommendation",
]
for index, label in enumerate(labels, start=3):
    cell = summary.cell(row=index, column=1, value=label)
    cell.font = Font(name="Arial", bold=True)
summary.column_dimensions["A"].width = 36
summary.column_dimensions["B"].width = 40

wb.save("templates/checkout_a11y_revenue_risk_template.xlsx")
print("template written")
